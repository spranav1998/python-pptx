"""Unit-test suite for animation support: `Slide.add_animation()` and related."""

from __future__ import annotations

import pytest

from pptx import Presentation as PresentationFactory
from pptx.enum.animation import PP_ANIMATION_TYPE
from pptx.oxml.ns import qn
from pptx.oxml.slide import (
    _ANIMATION_FILTER_MAP,
    _build_animation_click_par_xml,
    _build_animation_timing_xml,
)
from pptx.util import Inches


class DescribePP_ANIMATION_TYPE:
    """Unit-test suite for `PP_ANIMATION_TYPE` enumeration."""

    def it_has_the_expected_member_count(self):
        assert len(list(PP_ANIMATION_TYPE)) == 14

    def it_has_xml_values_for_all_members(self):
        for member in PP_ANIMATION_TYPE:
            assert member.xml_value is not None
            assert isinstance(member.xml_value, str)
            assert len(member.xml_value) > 0

    def it_maps_all_xml_values_to_the_filter_map(self):
        for member in PP_ANIMATION_TYPE:
            assert member.xml_value in _ANIMATION_FILTER_MAP

    def it_can_look_up_members_by_xml_value(self):
        assert PP_ANIMATION_TYPE.from_xml("fade") == PP_ANIMATION_TYPE.FADE
        assert PP_ANIMATION_TYPE.from_xml("appear") == PP_ANIMATION_TYPE.APPEAR


class Describe_build_animation_timing_xml:
    """Unit-test suite for `_build_animation_timing_xml()` helper."""

    def it_produces_valid_xml_for_appear(self):
        from pptx.oxml import parse_xml

        xml = _build_animation_timing_xml(2, "appear", 500)
        timing = parse_xml(xml)
        # Should have p:set but no p:animEffect
        sets = timing.findall(".//" + qn("p:set"))
        effects = timing.findall(".//" + qn("p:animEffect"))
        assert len(sets) == 1
        assert len(effects) == 0

    def it_produces_valid_xml_for_fade(self):
        from pptx.oxml import parse_xml

        xml = _build_animation_timing_xml(3, "fade", 1000)
        timing = parse_xml(xml)
        effects = timing.findall(".//" + qn("p:animEffect"))
        assert len(effects) == 1
        assert effects[0].get("filter") == "fade"
        assert effects[0].get("transition") == "in"

    def it_targets_the_correct_shape_id(self):
        from pptx.oxml import parse_xml

        xml = _build_animation_timing_xml(42, "fade", 500)
        timing = parse_xml(xml)
        spTgts = timing.findall(".//" + qn("p:spTgt"))
        assert any(t.get("spid") == "42" for t in spTgts)

    def it_sets_the_correct_duration(self):
        from pptx.oxml import parse_xml

        xml = _build_animation_timing_xml(2, "fade", 750)
        timing = parse_xml(xml)
        effects = timing.findall(".//" + qn("p:animEffect"))
        cTn = effects[0].find(".//" + qn("p:cTn"))
        assert cTn.get("dur") == "750"


class Describe_build_animation_click_par_xml:
    """Unit-test suite for `_build_animation_click_par_xml()` helper."""

    def it_produces_valid_xml_for_appear(self):
        from pptx.oxml import parse_xml

        xml = _build_animation_click_par_xml(5, "appear", 500, 7)
        par = parse_xml(xml)
        sets = par.findall(".//" + qn("p:set"))
        effects = par.findall(".//" + qn("p:animEffect"))
        assert len(sets) == 1
        assert len(effects) == 0

    def it_produces_valid_xml_for_wipe(self):
        from pptx.oxml import parse_xml

        xml = _build_animation_click_par_xml(5, "wipe_from_left", 800, 10)
        par = parse_xml(xml)
        effects = par.findall(".//" + qn("p:animEffect"))
        assert len(effects) == 1
        assert effects[0].get("filter") == "wipe(right)"

    def it_assigns_sequential_cTn_ids(self):
        from pptx.oxml import parse_xml

        xml = _build_animation_click_par_xml(5, "fade", 500, 7)
        par = parse_xml(xml)
        cTns = par.findall(".//" + qn("p:cTn"))
        ids = sorted([int(c.get("id")) for c in cTns])
        # Should start at 7 and be sequential
        assert ids[0] == 7


class DescribeSlide_add_animation:
    """Unit-test suite for `Slide.add_animation()` method."""

    def it_adds_a_fade_animation_to_a_shape(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        shape = slide.placeholders[0]
        shape.text = "Animated"

        slide.add_animation(shape, PP_ANIMATION_TYPE.FADE, duration=500)

        timing = slide._element.timing
        assert timing is not None
        effects = timing.findall(".//" + qn("p:animEffect"))
        assert len(effects) == 1
        assert effects[0].get("filter") == "fade"

    def it_adds_an_appear_animation_to_a_shape(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        shape = slide.placeholders[0]
        shape.text = "Appear"

        slide.add_animation(shape, PP_ANIMATION_TYPE.APPEAR)

        timing = slide._element.timing
        assert timing is not None
        sets = timing.findall(".//" + qn("p:set"))
        assert len(sets) == 1
        effects = timing.findall(".//" + qn("p:animEffect"))
        assert len(effects) == 0

    def it_can_add_multiple_animations_to_different_shapes(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.placeholders[0]
        title.text = "Title"
        body = slide.placeholders[1]
        body.text = "Body"

        slide.add_animation(title, PP_ANIMATION_TYPE.FADE, 500)
        slide.add_animation(body, PP_ANIMATION_TYPE.WIPE_FROM_LEFT, 800)

        timing = slide._element.timing
        effects = timing.findall(".//" + qn("p:animEffect"))
        assert len(effects) == 2
        filters = [e.get("filter") for e in effects]
        assert "fade" in filters
        assert "wipe(right)" in filters

    def it_targets_the_correct_shape(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        shape = slide.placeholders[0]
        shape.text = "Target"
        shape_id = shape.shape_id

        slide.add_animation(shape, PP_ANIMATION_TYPE.FADE, 500)

        timing = slide._element.timing
        spTgts = timing.findall(".//" + qn("p:spTgt"))
        assert any(t.get("spid") == str(shape_id) for t in spTgts)

    def it_uses_default_duration_of_500ms(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        shape = slide.placeholders[0]
        shape.text = "Default duration"

        slide.add_animation(shape, PP_ANIMATION_TYPE.FADE)

        timing = slide._element.timing
        effects = timing.findall(".//" + qn("p:animEffect"))
        cTn = effects[0].find(".//" + qn("p:cTn"))
        assert cTn.get("dur") == "500"

    def it_survives_a_save_load_roundtrip(self):
        import io
        import subprocess
        import sys

        # Run roundtrip in a subprocess to avoid test mock pollution
        code = """
import io, sys
from pptx import Presentation
from pptx.enum.animation import PP_ANIMATION_TYPE
from pptx.oxml.ns import qn

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.placeholders[0].text = "Roundtrip"
slide.placeholders[1].text = "Body"
slide.add_animation(slide.placeholders[0], PP_ANIMATION_TYPE.FADE, 1000)
slide.add_animation(slide.placeholders[1], PP_ANIMATION_TYPE.WIPE_FROM_BOTTOM, 750)

stream = io.BytesIO()
prs.save(stream)
stream.seek(0)

prs2 = Presentation(stream)
timing = prs2.slides[0]._element.timing
assert timing is not None
effects = timing.findall(".//" + qn("p:animEffect"))
assert len(effects) == 2, f"Expected 2 effects, got {len(effects)}"
print("ok")
"""
        result = subprocess.run(
            [sys.executable, "-c", code],
            capture_output=True,
            text=True,
        )
        assert result.returncode == 0, f"Subprocess failed: {result.stderr}"
        assert "ok" in result.stdout

    def it_raises_on_invalid_shape_type(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        with pytest.raises(TypeError, match="shape must be a BaseShape"):
            slide.add_animation("not a shape", PP_ANIMATION_TYPE.FADE)

    def it_raises_on_invalid_animation_type(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        shape = slide.placeholders[0]
        with pytest.raises(TypeError, match="animation_type must be"):
            slide.add_animation(shape, "fade")

    def it_raises_on_negative_duration(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        shape = slide.placeholders[0]
        with pytest.raises(ValueError, match="duration must be non-negative"):
            slide.add_animation(shape, PP_ANIMATION_TYPE.FADE, duration=-100)

    def it_works_with_all_animation_types(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        for i, anim_type in enumerate(PP_ANIMATION_TYPE):
            txbox = slide.shapes.add_textbox(
                Inches(0.5), Inches(0.3 * i), Inches(3), Inches(0.25)
            )
            txbox.text = anim_type.name
            slide.add_animation(txbox, anim_type, 500)

        timing = slide._element.timing
        assert timing is not None
        sets = timing.findall(".//" + qn("p:set"))
        # Each animation adds one p:set for visibility
        assert len(sets) == len(list(PP_ANIMATION_TYPE))
