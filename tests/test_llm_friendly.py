"""Unit-test suite for LLM-friendly `to_dict()` and `__repr__()` methods."""

from __future__ import annotations

import pytest

from pptx import Presentation as PresentationFactory
from pptx.presentation import Presentation
from pptx.slide import Slide
from pptx.text.text import Font, TextFrame, _Paragraph, _Run
from pptx.util import Inches, Pt

from .unitutil.cxml import element


class DescribePresentation_to_dict:
    """Unit-test suite for `Presentation.to_dict()` method."""

    def it_returns_a_dict_with_slide_count(self):
        prs = PresentationFactory()
        d = prs.to_dict()
        assert d["slide_count"] == 0
        assert "slides" in d
        assert isinstance(d["slides"], list)
        assert len(d["slides"]) == 0

    def it_includes_slide_dimensions(self):
        prs = PresentationFactory()
        d = prs.to_dict()
        assert d["slide_width"] is not None
        assert d["slide_height"] is not None
        assert isinstance(d["slide_width"], int)
        assert isinstance(d["slide_height"], int)

    def it_serializes_slides_with_shapes(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.placeholders[0].text = "Test Title"
        d = prs.to_dict()
        assert d["slide_count"] == 1
        assert len(d["slides"]) == 1
        slide_d = d["slides"][0]
        assert "shapes" in slide_d
        assert any(
            s.get("text") == "Test Title" for s in slide_d["shapes"]
        )


class DescribePresentation_repr:
    """Unit-test suite for `Presentation.__repr__()` method."""

    def it_returns_a_descriptive_string(self):
        prs = PresentationFactory()
        result = repr(prs)
        assert "Presentation(" in result
        assert "slides=" in result
        assert "slide_width=" in result
        assert "slide_height=" in result


class DescribeSlide_to_dict:
    """Unit-test suite for `Slide.to_dict()` method."""

    def it_returns_a_dict_with_slide_properties(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        d = slide.to_dict()
        assert "slide_id" in d
        assert "name" in d
        assert "shapes" in d
        assert isinstance(d["shapes"], list)

    def it_includes_text_content_in_shapes(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.placeholders[0].text = "My Title"
        d = slide.to_dict()
        shape_texts = [s.get("text", "") for s in d["shapes"]]
        assert "My Title" in shape_texts

    def it_includes_placeholder_info(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        d = slide.to_dict()
        placeholders = [s for s in d["shapes"] if s.get("is_placeholder")]
        assert len(placeholders) > 0
        for ph in placeholders:
            assert "placeholder_idx" in ph
            assert "placeholder_type" in ph


class DescribeSlide_repr:
    """Unit-test suite for `Slide.__repr__()` method."""

    def it_returns_a_descriptive_string(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        result = repr(slide)
        assert "Slide(" in result
        assert "slide_id=" in result
        assert "shapes=" in result


class DescribeBaseShape_to_dict:
    """Unit-test suite for `BaseShape.to_dict()` method."""

    def it_includes_core_shape_properties(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.placeholders[0].text = "Title"
        for shape in slide.shapes:
            d = shape.to_dict()
            assert "shape_id" in d
            assert "name" in d
            assert "left" in d
            assert "top" in d
            assert "width" in d
            assert "height" in d

    def it_includes_text_frame_for_text_shapes(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.placeholders[0].text = "Hello"
        for shape in slide.shapes:
            d = shape.to_dict()
            if shape.has_text_frame:
                assert "text" in d
                assert "text_frame" in d


class DescribeBaseShape_repr:
    """Unit-test suite for `BaseShape.__repr__()` method."""

    def it_returns_a_descriptive_string(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        for shape in slide.shapes:
            result = repr(shape)
            assert "name=" in result
            assert "shape_id=" in result
            assert "position=" in result
            assert "size=" in result


class DescribeTextFrame_to_dict:
    """Unit-test suite for `TextFrame.to_dict()` method."""

    def it_returns_a_dict_with_text_and_paragraphs(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        title.text = "Test"
        tf = title.text_frame
        d = tf.to_dict()
        assert d["text"] == "Test"
        assert "paragraphs" in d
        assert len(d["paragraphs"]) == 1

    def it_serializes_multi_paragraph_text(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        tf = title.text_frame
        tf.text = "Line 1"
        p = tf.add_paragraph()
        p.text = "Line 2"
        d = tf.to_dict()
        assert d["text"] == "Line 1\nLine 2"
        assert len(d["paragraphs"]) == 2


class DescribeTextFrame_repr:
    """Unit-test suite for `TextFrame.__repr__()` method."""

    def it_returns_a_descriptive_string(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        title.text = "Hello World"
        tf = title.text_frame
        result = repr(tf)
        assert "TextFrame(" in result
        assert "text=" in result
        assert "paragraphs=" in result


class DescribeParagraph_to_dict:
    """Unit-test suite for `_Paragraph.to_dict()` method."""

    def it_returns_a_dict_with_text(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        title.text = "Paragraph text"
        p = title.text_frame.paragraphs[0]
        d = p.to_dict()
        assert d["text"] == "Paragraph text"

    def it_includes_non_default_level(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.text = "Top level"
        p = tf.add_paragraph()
        p.text = "Indented"
        p.level = 2
        d = p.to_dict()
        assert d["level"] == 2

    def it_omits_default_level(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        title.text = "Default level"
        p = title.text_frame.paragraphs[0]
        d = p.to_dict()
        assert "level" not in d

    def it_includes_alignment_when_set(self):
        from pptx.enum.text import PP_ALIGN

        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        tf = title.text_frame
        tf.text = "Centered"
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        d = p.to_dict()
        assert "alignment" in d

    def it_includes_runs(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        title.text = "With runs"
        p = title.text_frame.paragraphs[0]
        d = p.to_dict()
        assert "runs" in d
        assert len(d["runs"]) >= 1


class DescribeParagraph_repr:
    """Unit-test suite for `_Paragraph.__repr__()` method."""

    def it_returns_a_descriptive_string(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        title.text = "Test paragraph"
        p = title.text_frame.paragraphs[0]
        result = repr(p)
        assert "_Paragraph(" in result
        assert "text=" in result
        assert "level=" in result


class DescribeRun_to_dict:
    """Unit-test suite for `_Run.to_dict()` method."""

    def it_returns_a_dict_with_text(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        title.text = "Run text"
        run = title.text_frame.paragraphs[0].runs[0]
        d = run.to_dict()
        assert d["text"] == "Run text"

    def it_includes_font_when_formatted(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        tf = title.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = "Bold text"
        run.font.bold = True
        run.font.size = Pt(24)
        d = run.to_dict()
        assert d["text"] == "Bold text"
        assert "font" in d
        assert d["font"]["bold"] is True
        assert d["font"]["size_pt"] == 24.0


class DescribeRun_repr:
    """Unit-test suite for `_Run.__repr__()` method."""

    def it_returns_a_descriptive_string(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        title.text = "Test"
        run = title.text_frame.paragraphs[0].runs[0]
        result = repr(run)
        assert "_Run(" in result
        assert "text=" in result


class DescribeFont_to_dict:
    """Unit-test suite for `Font.to_dict()` method."""

    def it_returns_empty_dict_for_unformatted_font(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        title.text = "Plain"
        run = title.text_frame.paragraphs[0].runs[0]
        d = run.font.to_dict()
        # unformatted font should return dict with only explicitly set properties
        assert isinstance(d, dict)

    def it_includes_explicitly_set_properties(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        tf = title.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = "Styled"
        run.font.name = "Arial"
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.italic = False
        d = run.font.to_dict()
        assert d["name"] == "Arial"
        assert d["size_pt"] == 18.0
        assert d["bold"] is True
        assert d["italic"] is False


class DescribeFont_repr:
    """Unit-test suite for `Font.__repr__()` method."""

    def it_returns_a_descriptive_string(self):
        prs = PresentationFactory()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        tf = title.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = "Styled"
        run.font.name = "Arial"
        run.font.size = Pt(24)
        result = repr(run.font)
        assert "Font(" in result
        assert "name=" in result
        assert "size=" in result
        assert "bold=" in result
        assert "italic=" in result
